import re
import shutil
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def update_tables_in_docx():
    # 1. Sao chép file gốc của bạn sang file tạm để tránh xung đột quyền đọc (read lock)
    try:
        shutil.copy('Tài_liệu_Kiến_trúc_Hệ_thống.docx', 'temp.docx')
        print("Copied to temp.docx")
    except Exception as e:
        print(f"Error copying file: {e}")
        return

    doc = Document('temp.docx')
    body = doc.element.body

    # 2. Tìm index của Mục 3 trong file XML của Word
    start_idx = None
    for idx, child in enumerate(body):
        if child.tag.endswith('p'):
            p_text = "".join(node.text for node in child.iter() if node.tag.endswith('t')).strip()
            # Khớp với "3. Workflow" hoặc "3." của file Word cũ của người dùng
            if p_text.startswith('3.'):
                start_idx = idx
                print(f"Found Section 3 at XML index {idx}")
                break

    # 3. Xóa các phần tử từ Mục 3 trở đi nhưng giữ lại sectPr ở cuối file
    if start_idx is not None:
        elements_to_delete = []
        for idx in range(start_idx, len(body)):
            child = body[idx]
            if not child.tag.endswith('sectPr'):
                elements_to_delete.append(child)
        for elem in elements_to_delete:
            body.remove(elem)
        print("Deleted old section 3 tables and content (preserved sectPr).")
    else:
        print("Could not find Section 3 in Word file.")
        return

    # 4. Đọc file system_documentation.md
    with open('system_documentation.md', 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')

    # Tìm dòng bắt đầu Mục 3 trong Markdown
    md_start_line = None
    for idx, line in enumerate(lines):
        if line.startswith('## 3.'):
            md_start_line = idx
            break

    if md_start_line is None:
        print("Không tìm thấy Mục 3 trong file system_documentation.md")
        return

    # 5. Đọc và chèn tiếp các phần từ Mục 3 trở xuống với định dạng bảng chuẩn của Word
    in_code_block = False
    i = md_start_line

    while i < len(lines):
        line = lines[i]
        
        # Đặc biệt: Nếu gặp khối mã mermaid, ta chèn ảnh sơ đồ vào
        if line.startswith('```mermaid'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                i += 1
            i += 1 # bỏ qua dấu đóng ```
            
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                from docx.shared import Inches
                run.add_picture('sequence_diagram.png', width=Inches(6.0))
                print("Successfully inserted sequence_diagram.png")
            except Exception as e:
                print(f"Error inserting sequence diagram: {e}")
            continue

        if line.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style.font.name = 'Courier New'
            p.style.font.size = Pt(10)
            i += 1
            continue
            
        # Kiểm tra và xử lý Bảng (Table)
        if line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 3:
                def split_row(r):
                    parts = [p.strip() for p in r.split('|')]
                    if parts[0] == '':
                        parts.pop(0)
                    if len(parts) > 0 and parts[-1] == '':
                        parts.pop()
                    return parts
                
                headers = split_row(table_lines[0])
                data_rows = [split_row(r) for r in table_lines[2:]]
                num_cols = len(headers)
                
                table = doc.add_table(rows=1 + len(data_rows), cols=num_cols)
                table.style = 'Table Grid'
                
                # Header
                hdr_cells = table.rows[0].cells
                for col_idx, text in enumerate(headers):
                    hdr_cells[col_idx].text = text
                    for paragraph in hdr_cells[col_idx].paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                
                # Data
                for row_idx, row_data in enumerate(data_rows):
                    row_cells = table.rows[row_idx + 1].cells
                    for col_idx, text in enumerate(row_data):
                        if col_idx < len(row_cells):
                            clean_text = re.sub(r'<br\s*/?>', '\n', text)
                            row_cells[col_idx].text = clean_text
            continue
            
        if line.startswith('# '):
            heading = doc.add_heading(line[2:], level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('---'):
            doc.add_page_break()
        elif line.startswith('* ') or line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line)
            doc.add_paragraph(text, style='List Number')
        else:
            if line.strip() != "":
                p = doc.add_paragraph()
                parts = re.split(r'(\*\*.*?\*\*|`.*?`)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    elif part.startswith('`') and part.endswith('`'):
                        run = p.add_run(part[1:-1])
                        run.font.name = 'Courier New'
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                    else:
                        p.add_run(part)
        i += 1

    # 6. Lưu đè lên file Word mới để không bị lỗi xung đột quyền ghi
    output_filename = 'Tai_Lieu_Kien_Truc_He_Thong_Word_Updated.docx'
    try:
        doc.save(output_filename)
        print(f"Successfully updated document: {output_filename}")
    except Exception as e:
        print(f"Error saving file: {e}")

if __name__ == '__main__':
    update_tables_in_docx()
