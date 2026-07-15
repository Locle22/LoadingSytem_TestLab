import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def md_to_docx(md_path, docx_path):
    doc = Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    # Split by lines
    lines = content.split('\n')
    
    in_code_block = False
    i = 0
    
    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style.font.name = 'Courier New'
            p.style.font.size = Pt(10)
            i += 1
            continue
            
        # Check for table
        if line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 3:
                def split_row(r):
                    parts = [p.strip() for p in r.split('|')]
                    if parts[0] == '':
                        parts.pop(0)
                    if len(parts) > 0 and parts[-1] == '':
                        parts.pop()
                    return parts
                
                headers = split_row(table_lines[0])
                data_rows = [split_row(r) for r in table_lines[2:]]
                num_cols = len(headers)
                
                table = doc.add_table(rows=1 + len(data_rows), cols=num_cols)
                table.style = 'Table Grid'
                
                # Write header
                hdr_cells = table.rows[0].cells
                for col_idx, text in enumerate(headers):
                    hdr_cells[col_idx].text = text
                    for paragraph in hdr_cells[col_idx].paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                
                # Write data
                for row_idx, row_data in enumerate(data_rows):
                    row_cells = table.rows[row_idx + 1].cells
                    for col_idx, text in enumerate(row_data):
                        if col_idx < len(row_cells):
                            clean_text = re.sub(r'<br\s*/?>', '\n', text)
                            row_cells[col_idx].text = clean_text
            continue
            
        if line.startswith('# '):
            heading = doc.add_heading(line[2:], level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('---'):
            doc.add_page_break()
        elif line.startswith('* ') or line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line)
            doc.add_paragraph(text, style='List Number')
        else:
            if line.strip() != "":
                p = doc.add_paragraph()
                parts = re.split(r'(\*\*.*?\*\*|`.*?`)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    elif part.startswith('`') and part.endswith('`'):
                        run = p.add_run(part[1:-1])
                        run.font.name = 'Courier New'
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                    else:
                        p.add_run(part)
        i += 1

    doc.save(docx_path)
    print("Convert success!")

md_to_docx('system_documentation.md', 'Tai_Lieu_Kien_Truc_He_Thong.docx')
