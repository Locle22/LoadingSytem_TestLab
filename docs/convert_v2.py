import re
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def md_to_docx(md_path, docx_path):
    doc = Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    # Làm sạch các ký tự LaTeX toán học đơn giản để hiển thị đẹp trong Word
    content = content.replace('$[-180^\\circ, 180^\\circ]$', '[-180°, 180°]')
    content = content.replace('$\\Delta$', 'Delta (Δ)')
    content = content.replace('$-180^\\circ$', '-180°')
    content = content.replace('$180^\\circ$', '180°')
    content = content.replace('$\\Delta = \\text{angle} - \\text{slot\\_angle}(\\text{slot})$', 'Delta = angle - slot_angle(slot)')

    lines = content.split('\n')
    
    in_code_block = False
    i = 0
    
    while i < len(lines):
        line = lines[i]
        
        # 1. Nhận diện Code block
        if line.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style.font.name = 'Courier New'
            p.style.font.size = Pt(9.5)
            # Thêm viền nhẹ hoặc màu chữ tối
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            i += 1
            continue
            
        # 2. Nhận diện hình ảnh: ![caption](img_path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', line.strip())
        if img_match:
            caption = img_match.group(1)
            img_path = img_match.group(2)
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(6.2))
                
                # Thêm nhãn hình ảnh dưới ảnh
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_cap = p_cap.add_run(f"Hình: {caption}")
                run_cap.font.size = Pt(9.5)
                run_cap.font.italic = True
                print(f"Successfully inserted image: {img_path}")
            except Exception as e:
                print(f"Error inserting image {img_path}: {e}")
            i += 1
            continue

        # 3. Nhận diện bảng (Table)
        if line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 3:
                def split_row(r):
                    parts = [p.strip() for p in r.split('|')]
                    if parts[0] == '':
                        parts.pop(0)
                    if len(parts) > 0 and parts[-1] == '':
                        parts.pop()
                    return parts
                
                headers = split_row(table_lines[0])
                data_rows = [split_row(r) for r in table_lines[2:]]
                num_cols = len(headers)
                
                table = doc.add_table(rows=1 + len(data_rows), cols=num_cols)
                table.style = 'Table Grid'
                
                # Header formatting
                hdr_cells = table.rows[0].cells
                for col_idx, text in enumerate(headers):
                    hdr_cells[col_idx].text = text
                    for paragraph in hdr_cells[col_idx].paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        for run in paragraph.runs:
                            run.bold = True
                            run.font.size = Pt(10)
                
                # Data rows formatting
                for row_idx, row_data in enumerate(data_rows):
                    row_cells = table.rows[row_idx + 1].cells
                    for col_idx, text in enumerate(row_data):
                        if col_idx < len(row_cells):
                            clean_text = re.sub(r'<br\s*/?>', '\n', text)
                            row_cells[col_idx].text = clean_text
                            # Format paragraphs inside cells
                            for paragraph in row_cells[col_idx].paragraphs:
                                paragraph.style.font.size = Pt(9.5)
            continue
            
        # 4. Nhận diện Tiêu đề (Headings)
        if line.startswith('# '):
            heading = doc.add_heading(line[2:], level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in heading.runs:
                run.font.size = Pt(18)
                run.bold = True
        elif line.startswith('## '):
            heading = doc.add_heading(line[3:], level=2)
            for run in heading.runs:
                run.font.size = Pt(14)
                run.bold = True
        elif line.startswith('### '):
            heading = doc.add_heading(line[4:], level=3)
            for run in heading.runs:
                run.font.size = Pt(11)
                run.bold = True
        elif line.startswith('---'):
            doc.add_page_break()
        elif line.startswith('* ') or line.startswith('- '):
            # Nhận diện inline styles trong list item
            p = doc.add_paragraph(style='List Bullet')
            parse_inline_styles(p, line[2:])
        elif re.match(r'^\d+\.\s', line):
            # Nhận diện inline styles trong list number
            text = re.sub(r'^\d+\.\s', '', line)
            p = doc.add_paragraph(style='List Number')
            parse_inline_styles(p, text)
        else:
            if line.strip() != "":
                p = doc.add_paragraph()
                parse_inline_styles(p, line)
        i += 1

    doc.save(docx_path)
    print(f"Created Word document: {docx_path}")

def parse_inline_styles(paragraph, text):
    # Nhận diện chữ đậm (**bold**) và mã nội dòng (`inline code`)
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.color.rgb = RGBColor(0xc7, 0x25, 0x4e) # Màu hồng code đặc trưng
        else:
            paragraph.add_run(part)

if __name__ == '__main__':
    # Chuyển Cwd sang thư mục chứa file
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    md_to_docx('system_documentation_v2.md', 'Tai_Lieu_Kien_Truc_He_Thong_V2.docx')
